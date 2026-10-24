import asyncio
import io
import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from docx import Document

from src.multimodal.extract import extract_images
from src.multimodal.provider import prepare_image, MultimodalProvider
from src.multimodal.store import ImageStore


def png(color="red"):
    out = io.BytesIO()
    Image.new("RGB", (128, 96), color).save(out, "PNG")
    return out.getvalue()


class FakeProvider:
    def __init__(self):
        self.models = {"multimodal_embedding": {"embed_dim": 3}, "embedding": {"embed_dim": 3}}
        self.failed = False
        self.calls = 0
        self.version = 1

    def signature(self):
        return {"version": self.version}

    async def describe(self, image):
        self.calls += 1
        if self.failed:
            raise ValueError("model offline")
        return {"ocr": "", "description": "red"}

    async def embed(self, **kwargs):
        return [1, 0, 0]

    async def verify_evidence(self, question, query_image, candidates):
        return {'query_summary': 'red', 'matches': [
            {'id': key, 'verdict': 'same_image' if query_image is not None else 'relevant', 'reason': 'matching test fixture'}
            for key, _ in candidates]}


def test_decode_limits_and_bad_formats():
    assert prepare_image(png())[2:] == (128, 96)
    with pytest.raises(Exception):
        prepare_image(b"not an image")
    with pytest.raises(ValueError):
        prepare_image(b"x" * (20 * 1024 * 1024 + 1))


def test_pdf_raster_and_vector_provenance(tmp_path):
    path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((40, 40), "Deployment architecture")
    page.insert_image(fitz.Rect(40, 70, 180, 170), stream=png())
    page.draw_rect(fitz.Rect(200, 80, 300, 120))
    doc.new_page().insert_image(fitz.Rect(0, 0, 100, 100), stream=png())
    doc.save(path); doc.close()
    records = list(extract_images(path))
    assert len(records) == 4
    assert records[0][1]["page"] == 1
    assert "Deployment" in records[0][2]
    assert sum(x[1]["kind"] == "page_image" for x in records) == 2


def test_docx_table_and_repeated_picture(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Architecture", 1)
    doc.add_paragraph("Related source text")
    doc.add_picture(io.BytesIO(png()))
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(png()))
    doc.save(path)
    records = list(extract_images(path))
    assert len(records) == 2
    assert records[0][0] == records[1][0]
    assert records[0][1]["paragraph"] != records[1][1]["paragraph"]
    assert all("page" not in x[1] for x in records)


def test_shared_image_vectors_and_document_deletion(tmp_path):
    async def run():
        store = ImageStore(tmp_path / "store")
        path = tmp_path / "image.png"; path.write_bytes(png())
        provider = FakeProvider()
        await store.index("a", "one.png", path, provider)
        await store.index("b", "two.png", path, provider)
        assert provider.calls == 1
        hits = await store.search(provider, text="red")
        assert len(hits) == 1 and len(hits[0]["sources"]) == 2
        image_id = hits[0]["image_id"]
        store.remove_document("a")
        assert store.asset_path(image_id).exists()
        assert len((await store.search(provider, image=png()))[0]["sources"]) == 1
        store.remove_document("b")
        assert not store.asset_path(image_id).exists()
        assert await store.search(provider, text="red") == []
    asyncio.run(run())


def test_failed_rebuild_preserves_generation_and_reports_partial(tmp_path):
    async def run():
        store = ImageStore(tmp_path / "store"); provider = FakeProvider()
        path = tmp_path / "image.png"; path.write_bytes(png())
        await store.index("a", path.name, path, provider)
        old = store.occurrences()[0]["image_id"]
        path.write_bytes(png("blue")); provider.failed = True
        with pytest.raises(ValueError, match="部分完成"):
            await store.index("a", path.name, path, provider)
        assert store.occurrences()[0]["image_id"] == old
        assert store.list_documents()[0]["status"] == "partial"
        provider.failed = False
        await store.index("a", path.name, path, provider)
        assert store.occurrences()[0]["image_id"] != old
    asyncio.run(run())


def test_cancel_and_changed_document_never_publish(tmp_path):
    async def run():
        store = ImageStore(tmp_path / "store"); provider = FakeProvider()
        path = tmp_path / "image.png"; path.write_bytes(png())
        with pytest.raises(asyncio.CancelledError):
            await store.index("a", path.name, path, provider, cancelled=lambda: True)
        assert not store.occurrences()
        with pytest.raises(ValueError, match="已删除"):
            await store.index("a", path.name, path, provider, valid=lambda: False)
        assert not store.occurrences()
    asyncio.run(run())


def test_signature_mismatch_and_source_invalidation(tmp_path):
    async def run():
        store = ImageStore(tmp_path / "store"); provider = FakeProvider()
        path = tmp_path / "image.png"; path.write_bytes(png())
        await store.index("a", path.name, path, provider)
        provider.version = 2
        with pytest.raises(ValueError, match="重建"):
            await store.search(provider, text="red")
        store.invalidate("a")
        assert not store.occurrences()
        assert await store.search(provider, text="red") == []
    asyncio.run(run())


def test_attachments_workspace_session_and_expiry(tmp_path):
    first = ImageStore(tmp_path / "one"); second = ImageStore(tmp_path / "two")
    key = first.attachment(png())
    with pytest.raises(ValueError): second.attachment_path(key)
    first.bind_attachment(key, "session1")
    with pytest.raises(ValueError): first.attachment_path(key, "session2")
    assert first.attachment_path(key, "session1").exists()
    first.expire_attachments("session1")
    with pytest.raises(ValueError): first.attachment_path(key, "session1")
    key = first.attachment(png())
    with first.db() as db: db.execute("UPDATE attachments SET created=0")
    with pytest.raises(ValueError): first.attachment_path(key)
    first.expire_attachments()
    assert not (first.root / "attachments" / f"{key}.jpg").exists()


def test_clear_preserves_uploaded_file(tmp_path):
    async def run():
        store = ImageStore(tmp_path / "store"); path = tmp_path / "source.png"; path.write_bytes(png())
        await store.index("a", path.name, path, FakeProvider())
        store.clear()
        assert path.exists()
        assert not store.occurrences()
    asyncio.run(run())


def test_chat_request_accepts_image_only_and_rejects_empty():
    from src.api.server import ChatSendRequest
    assert ChatSendRequest(image_id="a"*32).message == ""
    with pytest.raises(ValueError): ChatSendRequest(message=" ")


def test_provider_wire_format_and_bad_dimensions(monkeypatch):
    import httpx
    real_client = httpx.AsyncClient
    requests = []
    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(200, json={"data": [{"embedding": [1, 0, 0]}]})
    monkeypatch.setattr(httpx, "AsyncClient", lambda **kwargs: real_client(transport=httpx.MockTransport(respond)))
    model = {"model": "vl", "base_url": "https://example.com/v1", "api_key": "test", "adapter": "siliconflow_vl", "embed_dim": 3}
    p = MultimodalProvider({"multimodal_embedding": model})
    async def run():
        await p.embed(image=png())
        await p.embed(text="red")
        assert "image" in requests[0]["input"]
        assert requests[1]["input"] == {"text": "red"}
        model["embed_dim"] = 4
        with pytest.raises(ValueError, match="维度"):
            await p.embed(text="red")
    asyncio.run(run())


def test_description_recovers_from_disconnect_and_invalid_json(monkeypatch):
    import httpx
    from src.multimodal import provider
    async def no_sleep(_): pass
    monkeypatch.setattr(provider.asyncio, 'sleep', no_sleep)
    p = MultimodalProvider({})
    budgets = []
    async def stream(messages, max_tokens):
        budgets.append(max_tokens)
        if len(budgets) == 1:
            raise httpx.RemoteProtocolError('disconnected')
        yield '{"ocr":' if len(budgets) == 2 else '{"ocr":"ready","description":"log"}'
    monkeypatch.setattr(p, 'stream', stream)
    assert asyncio.run(p.describe(png()))['ocr'] == 'ready'
    assert budgets == [4096, 8192, 8192]


@pytest.mark.parametrize('status,attempts', [(503, 3), (401, 1)])
def test_embedding_retries_only_transient_http_errors(monkeypatch, status, attempts):
    import httpx
    from src.multimodal import provider
    async def no_sleep(_): pass
    monkeypatch.setattr(provider.asyncio, 'sleep', no_sleep)
    real_client = httpx.AsyncClient
    requests = []
    def respond(request):
        requests.append(request)
        return httpx.Response(status)
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kw: real_client(transport=httpx.MockTransport(respond)))
    p = MultimodalProvider({'multimodal_embedding': {'model': 'vl', 'base_url': 'https://example.com', 'api_key': 'test', 'adapter': 'siliconflow_vl'}})
    with pytest.raises(ValueError, match=str(status)):
        asyncio.run(p.embed(text='test'))
    assert len(requests) == attempts


def test_vision_length_finish_is_not_treated_as_complete(monkeypatch):
    import httpx
    from src.multimodal.provider import RetryableModelError
    real_client = httpx.AsyncClient
    def respond(request):
        return httpx.Response(200, text='data: {"choices":[{"delta":{"content":"partial"},"finish_reason":null}]}\n\ndata: {"choices":[{"delta":{},"finish_reason":"length"}]}\n\ndata: [DONE]\n\n')
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kw: real_client(transport=httpx.MockTransport(respond)))
    p = MultimodalProvider({'vision': {'model': 'vision', 'base_url': 'https://example.com', 'api_key': 'test'}})
    async def run():
        return [part async for part in p.stream([], 10)]
    with pytest.raises(RetryableModelError, match='token 上限'):
        asyncio.run(run())


def test_image_only_manifest_publication_exemption_is_narrow(tmp_path):
    from src.index_validation import validate_index
    item = {"doc_name": "a.png", "file_type": "image", "char_count": 0, "text_index_exempt": True}
    assert validate_index(tmp_path, {"documents": {"a": item}}, ["a.png"])["chunks"] == 0
    item["char_count"] = 10
    with pytest.raises(RuntimeError, match="exemption"):
        validate_index(tmp_path, {"documents": {"a": item}}, ["a.png"])


def test_original_citation_does_not_resolve_to_another_document(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from src.api import server
    monkeypatch.setattr(server, "get_config", lambda: {"paths": {"data_dir": str(tmp_path)}})
    images = server._image_store("one")
    path = tmp_path / "sample.png"; path.write_bytes(png())
    asyncio.run(images.index("a", path.name, path, FakeProvider()))
    asyncio.run(images.index("b", path.name, path, FakeProvider()))
    key = images.occurrences()[0]["image_id"]
    client = TestClient(server.app)
    assert client.get(f"/api/kb/images/{key}?workspace=one&doc_id=a").status_code == 200
    assert client.get(f"/api/kb/images/{key}?workspace=two").status_code == 404
    images.remove_document("a")
    assert client.get(f"/api/kb/images/{key}?workspace=one&doc_id=a").status_code == 404
    assert client.get(f"/api/kb/images/{key}?workspace=one&doc_id=b").status_code == 200
    assert client.get("/api/kb/image-documents?workspace=../../outside").status_code == 422


@pytest.mark.parametrize('index_busy', [False, True])
@pytest.mark.parametrize('question', ['', '这张图片出自哪个文档'])
@pytest.mark.parametrize('accept_evidence', [False, True])
def test_vision_stream_and_persisted_attachment(tmp_path, monkeypatch, index_busy, question, accept_evidence):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    from src.api import server
    from src.multimodal import api
    monkeypatch.setattr(server, "get_config", lambda: {"paths": {"data_dir": str(tmp_path)}})
    monkeypatch.setattr(server, "_ensure_workspace_available", lambda _: None)
    images = server._image_store("one")
    path = tmp_path / "sample.png"; path.write_bytes(png())
    fake = FakeProvider()
    asyncio.run(images.index("a", path.name, path, fake))
    if not accept_evidence:
        async def uncertain(*args):
            return {'query_summary': '', 'matches': []}
        fake.verify_evidence = uncertain
    sent = []
    async def stream(messages, max_tokens):
        sent.extend(messages)
        yield "有相关图片"
        yield "[1]。"
    fake.stream = stream
    monkeypatch.setattr(api, "provider", lambda _: fake)
    class Service:
        async def preview_context(self, *args, **kwargs):
            assert not question, 'Source lookup must not retrieve generic text'
            return {"data": {}}
        def _citations_from_raw(self, raw): return []
    monkeypatch.setattr(server, "get_lightrag_service", lambda _: Service())
    sid = "a"*12
    session = {"id": sid, "workspace": "one", "title": "新对话", "messages": []}
    monkeypatch.setattr(server, "_ensure_session", lambda *args: (sid, session))
    monkeypatch.setattr(server, "_load_session", lambda _: session)
    monkeypatch.setattr(server, "_save_session", lambda *args: None)
    settings = server.ChatSettings(context_window=32768)
    monkeypatch.setattr(server, "_request_chat_settings", lambda *args: settings)
    key = images.attachment(png())
    workspace_lock = asyncio.Lock()
    if index_busy:
        asyncio.run(workspace_lock.acquire())
    monkeypatch.setattr(server, '_get_workspace_rag_lock', lambda _: workspace_lock)
    client = TestClient(server.app)
    response = client.post("/api/chat/send/stream", json={"workspace": "one", "session_id": sid, "image_id": key, "message": question})
    assert response.status_code == 200
    assert "event: done" in response.text and "event: error" not in response.text
    assert ('"content_type": "image"' in response.text) == accept_evidence
    assert session["messages"][0]["image_id"] == key
    if accept_evidence:
        assert session["messages"][1]["citations"][0]["doc_id"] == "a"
    else:
        assert session['messages'][1]['citations'] == []
    if question:
        assert '文本检索不可用' not in response.text
        assert not sent  # Source names are rendered from verified metadata, not generated.
        assert ('来源如下' in response.text) == accept_evidence
        if not accept_evidence:
            assert '未能确认' in response.text
    else:
        parts = sent[-1]['content']
        assert len([p for p in parts if p['type'] == 'image_url']) == 1
    assert client.get(f"/api/images/attachments/{key}?workspace=one&session_id={sid}").status_code == 200
    assert client.get(f"/api/images/attachments/{key}?workspace=one").status_code == 404
    assert not server._session_turn_locks.get(sid, asyncio.Lock()).locked()
    if index_busy:
        assert workspace_lock.locked()  # Retrieval must not release another task's lock.
        workspace_lock.release()
    settings.context_window = 1024
    response = client.post("/api/chat/send/stream", json={"workspace": "one", "session_id": sid, "image_id": key})
    assert "event: error" in response.text and "预算不足" in response.text


@pytest.mark.parametrize('question', ['这张图片出自哪个文档', '帮我看看', '有没有故障', '解释这张图', 'what is happening here'])
def test_source_lookup_uses_image_channel_without_generic_text_votes(tmp_path, question):
    async def run():
        store = ImageStore(tmp_path / 'images')
        path = tmp_path / 'red.png'
        path.write_bytes(png())
        provider = FakeProvider()
        await store.index('doc', path.name, path, provider)
        calls = []
        async def embed(**kwargs):
            calls.append(kwargs)
            return [1, 0, 0]
        provider.embed = embed
        hits = await store.search(provider, text=question, image=png())
        assert len(hits) == 1
        assert len(calls) == 1 and 'image' in calls[0]
        calls.clear()
        await store.search(provider, text='红色背景', image=png())
        assert len(calls) == 1  # Every attachment uses the image as the retrieval anchor.
    asyncio.run(run())


def test_chat_saves_before_retrieval_and_survives_viewer_disconnect(monkeypatch):
    from types import SimpleNamespace
    from datetime import datetime, timezone
    from src.multimodal import api
    from fastapi.responses import StreamingResponse
    async def run():
        session = {'messages': [], 'title': '新对话'}
        turn_lock, session_lock = asyncio.Lock(), asyncio.Lock()
        await turn_lock.acquire()
        ready, finish = asyncio.Event(), asyncio.Event()
        tasks = []
        def spawn(coro):
            tasks.append(asyncio.create_task(coro))
        s = SimpleNamespace(
            datetime=datetime, timezone=timezone,
            _ensure_workspace_available=lambda _: None,
            _ensure_session=lambda *args: ('session', session),
            _get_session_lock=lambda _: session_lock,
            _load_session=lambda _: session,
            _save_session=lambda *args: None,
            _spawn_background=spawn,
            _release_session_turn_lock=lambda sid, lock: lock.release(),
        )
        async def impl(s, req, lock, *, turn_index, status_callback=None):
            ready.set()
            await finish.wait()
            async def events():
                session['messages'][turn_index+1].update(content='finished', generation_status='complete')
                yield 'event: done\ndata: {}\n\n'
            return StreamingResponse(events())
        monkeypatch.setattr(api, '_chat_impl', impl)
        req = SimpleNamespace(session_id='session', workspace='default', message='question', image_id=None)
        response = await api.chat(s, req, turn_lock)
        assert session['messages'][0]['content'] == 'question'
        assert session['messages'][1]['generation_status'] == 'running'
        assert 'event: session' in await anext(response.body_iterator)
        await ready.wait()
        await response.body_iterator.aclose()
        assert not tasks[0].done()
        finish.set()
        await tasks[0]
        assert len(session['messages']) == 2
        assert session['messages'][1]['content'] == 'finished'
        assert not turn_lock.locked()
    asyncio.run(run())


def test_restart_marks_pending_turn_interrupted_without_losing_question(tmp_path, monkeypatch):
    from src.api import server
    sid = 'f' * 12
    (tmp_path / f'{sid}.json').write_text('{}')
    session = {'messages': [
        {'role': 'user', 'content': 'question'},
        {'role': 'assistant', 'content': 'pending', 'generation_status': 'running'},
    ]}
    saved = []
    monkeypatch.setattr(server, 'SESSIONS_DIR', tmp_path)
    monkeypatch.setattr(server, '_load_session', lambda _: session)
    monkeypatch.setattr(server, '_save_session', lambda *args: saved.append(args))
    server._recover_interrupted_chat_turns()
    assert session['messages'][0]['content'] == 'question'
    assert session['messages'][1]['generation_status'] == 'failed'
    assert len(saved) == 1


def test_long_ocr_indexes_all_parts(tmp_path):
    class LongProvider(FakeProvider):
        async def describe(self, image): return {"ocr": "甲"*900 + "末尾内容", "description": "示例"}
        async def embed(self, **kwargs):
            if not kwargs.get("visual", True): texts.append(kwargs["text"])
            return [1, 0, 0]
    texts = []
    store = ImageStore(tmp_path / "store"); path = tmp_path / "sample.png"; path.write_bytes(png())
    asyncio.run(store.index("a", path.name, path, LongProvider()))
    assert "末尾内容" in "".join(texts)
