"""Opt-in real API smoke test; all test documents live in an isolated data root.

Start: python scripts/multimodal_smoke.py serve
Verify: python scripts/multimodal_smoke.py verify
Uses existing encrypted model credentials without printing them.
"""
from __future__ import annotations

import copy
import io
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUTPUT = ROOT / ".workbuddy" / "multimodal-validation"


def serve():
    import uvicorn
    from src import config_loader
    from src.model_profiles import _load_store, _load_keys, _save_store, _save_keys
    original = config_loader.get_config()
    profiles, keys = _load_store(original), _load_keys(original)
    config = copy.deepcopy(original)
    data = OUTPUT / "data"
    config["paths"] = {"data_dir": str(data), "lightrag_dir": str(data / "lightrag"),
                       "lightrag_manifest": str(data / "lightrag_manifest.json"), "docs_dir": str(OUTPUT / "fixtures")}
    config_loader._config = config
    _save_store(profiles, config); _save_keys(keys, config)
    from src.api.server import app
    uvicorn.run(app, host="127.0.0.1", port=8101)


def verify():
    import httpx
    import fitz
    from PIL import Image, ImageDraw
    from docx import Document
    fixtures = OUTPUT / "fixtures"
    fixtures.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 130, 230, 230), fill="#dbeafe", outline="#2563eb", width=3)
    draw.text((85, 170), "NGINX PROXY", fill="black")
    draw.line((230, 180, 370, 100), fill="black", width=3)
    draw.line((230, 180, 370, 270), fill="black", width=3)
    for y, label in ((60, "APP SERVER A"), (230, "APP SERVER B")):
        draw.rectangle((370, y, 580, y+80), fill="#dcfce7", outline="#16a34a", width=3)
        draw.text((400, y+35), label, fill="black")
    image.save(fixtures / "architecture.png")
    blue = Image.new("RGB", (320, 240), "#172554")
    blue.save(fixtures / "blue.png")
    pdf = fitz.open(); page = pdf.new_page()
    page.insert_text((40, 40), "Deployment: Nginx forwards traffic to two application servers.")
    page.insert_image(fitz.Rect(40, 80, 550, 370), filename=str(fixtures / "architecture.png"))
    pdf.save(fixtures / "deployment.pdf"); pdf.close()
    word = Document(); word.add_heading("Deployment", 1)
    word.add_paragraph("Nginx forwards traffic to application servers A and B.")
    word.add_picture(str(fixtures / "architecture.png")); word.save(fixtures / "deployment.docx")
    report = {"fixture_type": "synthetic functional fixtures, not retrieval quality benchmark", "steps": []}
    with httpx.Client(base_url="http://127.0.0.1:8101", timeout=180) as client:
        def checked(response):
            response.raise_for_status(); return response.json()
        names = ["architecture.png", "blue.png", "deployment.pdf", "deployment.docx"]
        for name in names:
            result = checked(client.post("/api/kb/upload", files={"file": (name, (fixtures / name).read_bytes())}))
            report["steps"].append({"upload": name, "doc_id": result["doc_id"]})
        started = time.perf_counter()
        task = checked(client.post("/api/kb/index-images", json={"doc_names": names}))
        while True:
            task = checked(client.get("/api/kb/index-tasks/" + task["task_id"]))
            if task["status"] not in {"queued", "running"}: break
            if time.perf_counter() - started > 600: raise TimeoutError("image indexing exceeded 10 minutes")
            time.sleep(2)
        report["index"] = {"status": task["status"], "seconds": time.perf_counter()-started, "results": task["results"], "errors": task["errors"]}
        assert task["status"] == "succeeded", report["index"]
        started = time.perf_counter()
        hits = checked(client.post("/api/recall/images", json={"query": "Nginx proxy with two application servers"}))["hits"]
        report["text_search"] = {"seconds": time.perf_counter()-started, "sources": [h["sources"] for h in hits]}
        assert hits and hits[0]["description"]
        attachment = checked(client.post("/api/images/attachments", files={"file": ("query.png", (fixtures / "architecture.png").read_bytes())}))
        hits = checked(client.post("/api/recall/images", json={"image_id": attachment["image_id"]}))["hits"]
        assert any(s["doc_name"] in {"architecture.png", "deployment.pdf", "deployment.docx"} for s in hits[0]["sources"])
        report["image_search"] = {"top_sources": hits[0]["sources"]}
        result = client.post("/api/chat/send/stream", json={"image_id": attachment["image_id"], "message": "请根据检索到的资料说明这张架构图。"})
        result.raise_for_status()
        events = result.text
        assert "event: error" not in events, events
        assert "event: done" in events and '"content_type": "image"' in events
        report["chat_sse"] = events
        OUTPUT.mkdir(parents=True, exist_ok=True)
        (OUTPUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"ok": True, "index_seconds": report["index"]["seconds"], "text_search_seconds": report["text_search"]["seconds"], "report": str(OUTPUT / "report.json")}, ensure_ascii=False))


if __name__ == "__main__":
    {"serve": serve, "verify": verify}[sys.argv[1]]()
